import os
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Explicitly target your apartment CSV data file
input_file = "riyadh_cleaned_apartments.csv"
output_file = "riyadh_cleaned_apartments_formatted.xlsx"

print(f"📂 Processing file: {input_file}")

# Load the CSV file safely
df = pd.read_csv(input_file, on_bad_lines='skip')

# Clean up trailing spaces in text columns
for col in df.select_dtypes(include=['object']).columns:
    df[col] = df[col].astype(str).str.strip().str.replace(r'\s+', ' ', regex=True)

# Save to an Excel workbook
with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
    df.to_excel(writer, index=False, sheet_name='Apartments')

# Apply professional formatting to fix spacing and alignment
wb = openpyxl.load_workbook(output_file)
ws = wb.active

# Show gridlines for a clean spreadsheet grid look
ws.views.sheetView[0].showGridLines = True

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
regular_font = Font(name="Calibri", size=11)
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")

thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

# Format header row
for col_num in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = align_center
    cell.border = thin_border

# Format data rows and adjust column widths cleanly
for col_num in range(1, ws.max_column + 1):
    col_letter = get_column_letter(col_num)
    max_len = 0
    
    for row_num in range(2, ws.max_row + 1):
        cell = ws.cell(row=row_num, column=col_num)
        cell.font = regular_font
        cell.border = thin_border
        
        if isinstance(cell.value, (int, float)):
            cell.alignment = align_right
            if cell.value > 1000:
                cell.number_format = '#,##0'
        else:
            cell.alignment = align_left
            
        if cell.value:
            cell_len = len(str(cell.value))
            if cell_len > max_len:
                max_len = cell_len
                
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 35)

wb.save(output_file)
print(f"✨ Successfully formatted and saved as: {output_file}")